import io
import os
import re
import tempfile
import zipfile
import logging
from pathlib import Path
from xml.etree import ElementTree as ET

from flask import Flask, request, send_file, jsonify
from flask_cors import CORS
from pptx import Presentation
from pptx.util import Inches

app = Flask(__name__)
CORS(app)
app.config["MAX_CONTENT_LENGTH"] = 100 * 1024 * 1024

logging.basicConfig(level=logging.INFO)
log = logging.getLogger("cbtrade-converter")

CONTROL_RE = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


def clean_xml_bytes(data: bytes) -> bytes:
    text = data.decode("utf-8")
    text = CONTROL_RE.sub("", text)
    root = ET.fromstring(text)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def clean_pptx(src: bytes, logo_bytes: bytes | None = None) -> bytes:
    with tempfile.TemporaryDirectory(prefix="cbtrade_") as td:
        src_path = Path(td) / "input.pptx"
        out_path = Path(td) / "output.pptx"
        src_path.write_bytes(src)

        cleaned = Path(td) / "cleaned.pptx"
        with zipfile.ZipFile(src_path, "r") as zin, zipfile.ZipFile(cleaned, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                payload = zin.read(info.filename)
                if info.filename.endswith(".xml") or info.filename.endswith(".rels"):
                    try:
                        payload = clean_xml_bytes(payload)
                    except Exception:
                        pass
                zout.writestr(info, payload)

        if logo_bytes:
            prs = Presentation(str(cleaned))
            logo_path = Path(td) / "logo.png"
            logo_path.write_bytes(logo_bytes)
            for slide in prs.slides:
                slide.shapes.add_picture(
                    str(logo_path),
                    prs.slide_width - Inches(1.55),
                    Inches(0.18),
                    width=Inches(1.25),
                )
            prs.save(str(out_path))
        else:
            out_path.write_bytes(cleaned.read_bytes())

        return out_path.read_bytes()


@app.get("/")
def index():
    return jsonify({"service": "承邦有限公司一鍵轉檔後端", "status": "ok", "endpoint": "/convert"})


@app.get("/health")
def health():
    return jsonify({"status": "ok"})


@app.post("/convert")
def convert():
    uploaded = request.files.get("file")
    logo = request.files.get("logo")

    if not uploaded or not uploaded.filename:
        return jsonify({"error": "請上傳 PPTX 或 PDF 檔案"}), 400

    filename = uploaded.filename
    suffix = Path(filename).suffix.lower()
    if suffix not in {".pptx", ".pdf"}:
        return jsonify({"error": "只支援 PPTX 或 PDF"}), 400

    try:
        data = uploaded.read()
        if not data:
            return jsonify({"error": "上傳檔案是空的"}), 400

        if suffix == ".pptx":
            logo_bytes = logo.read() if logo and logo.filename else None
            result = clean_pptx(data, logo_bytes)
            return send_file(
                io.BytesIO(result),
                mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                as_attachment=True,
                download_name=f"承邦轉檔_{Path(filename).name}",
            )

        return send_file(
            io.BytesIO(data),
            mimetype="application/pdf",
            as_attachment=True,
            download_name=f"承邦轉檔_{Path(filename).name}",
        )

    except zipfile.BadZipFile:
        return jsonify({"error": "PPTX 檔案結構損壞或不是有效的 PPTX"}), 422
    except Exception as exc:
        log.exception("Conversion failed for %s", filename)
        return jsonify({"error": f"轉檔失敗：{type(exc).__name__}: {exc}"}), 500


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "10000")))
